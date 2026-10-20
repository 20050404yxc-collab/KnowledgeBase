# 文件：check_python.py
# 层：维护脚本
# 作用：准备或验证开发环境。
# 说明：使用隔离资源，不代表正式接口验收。
import hashlib
import importlib.metadata
import io
import json
import os
import sqlite3
import tempfile
import time
import uuid
from pathlib import Path

import httpx
import keyring
import numpy as np
import onnxruntime
import pdfplumber
import pypdfium2
import sqlite_vec
from docx import Document
from markdown_it import MarkdownIt
from PIL import Image, ImageDraw, ImageFont
from pydantic import BaseModel
from rapidocr import RapidOCR
from rapidocr.utils.download_file import DownloadFile


# 构造可丢弃的文字 PDF，仅用于本地环境验证。
def synthetic_pdf():
    content = b"BT /F1 20 Tf 60 150 Td (Knowledge 123) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 500 200] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(content)} >>\nstream\n".encode() + content + b"\nendstream",
    ]
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(data)
    data.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(data)


# 用合成对象验证结构校验、模拟 HTTP、Markdown 与 DOCX。
def check_data_libraries() -> list[str]:
    checks = []

    class Payload(BaseModel):
        ready: bool

    assert Payload.model_validate_json('{"ready":true}').ready
    with httpx.Client(
        transport=httpx.MockTransport(lambda _: httpx.Response(200, json={"ready": True}))
    ) as client:
        assert client.get("https://environment.invalid/check").json()["ready"]
    checks.append("Pydantic validation and httpx mock transport")

    markdown = MarkdownIt().render("# 环境检查\n\nKnowledge 123")
    assert "<h1>环境检查</h1>" in markdown
    document = Document()
    document.add_heading("环境检查", 1)
    document.add_paragraph("Knowledge 123")
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    assert Document(buffer).paragraphs[1].text == "Knowledge 123"
    checks.append("Markdown and DOCX synthetic document roundtrip")

    return checks


# 用合成 PDF 验证文字和页面渲染，及时关闭页面及位图。
def check_pdf() -> str:
    pdf = synthetic_pdf()
    with pdfplumber.open(io.BytesIO(pdf)) as document:
        assert document.pages[0].extract_text() == "Knowledge 123"
    with pypdfium2.PdfDocument(pdf) as document:
        page = document[0]
        bitmap = page.render(scale=1)
        assert bitmap.width == 500 and bitmap.height == 200
        bitmap.close()
        page.close()
    return "PDF text extraction and PDFium page rendering"


# 验证内存 SQLite 扩展，加载后立即关闭扩展加载权限。
def check_sqlite() -> str:
    with sqlite3.connect(":memory:") as connection:
        connection.enable_load_extension(True)
        try:
            sqlite_vec.load(connection)
        finally:
            connection.enable_load_extension(False)
        assert connection.execute("SELECT vec_version()").fetchone()[0] == "v0.1.9"
        assert connection.execute("SELECT vec_length('[1,2,3]')").fetchone()[0] == 3
        connection.execute("CREATE VIRTUAL TABLE documents USING fts5(content)")
        connection.execute("INSERT INTO documents VALUES ('knowledge environment')")
        assert (
            connection.execute(
                "SELECT count(*) FROM documents WHERE documents MATCH 'knowledge'"
            ).fetchone()[0]
            == 1
        )
    return "SQLite FTS5 and sqlite-vec Windows extension (English fixture)"


# 使用唯一合成服务名验证系统凭据，并保证删除测试凭据。
def check_credentials() -> str:
    credential_service = "KnowledgeBase.EnvironmentCheck." + str(uuid.uuid4())
    try:
        keyring.set_password(credential_service, "synthetic", "environment-check")
        assert keyring.get_password(credential_service, "synthetic") == "environment-check"
    finally:
        keyring.delete_password(credential_service, "synthetic")
    return "Windows credential synthetic write/read/delete"


# 只使用已校验本地模型执行合成图片检查，禁止隐式下载。
def check_ocr() -> tuple[str, float]:
    model_root = Path(os.environ.get("KB_OCR_MODEL_DIR", "D:/dev-models/KnowledgeBase/rapidocr"))
    receipt = json.loads((model_root / "models.lock.json").read_text(encoding="utf-8"))
    for model in receipt["models"]:
        with (model_root / model["file"]).open("rb") as stream:
            assert hashlib.file_digest(stream, "sha256").hexdigest() == model["sha256"]
    assert "CPUExecutionProvider" in onnxruntime.get_available_providers()

    # 环境检查禁止下载模型，缺失资源应明确失败。
    def disallow_download(*args, **kwargs):
        raise RuntimeError("Smoke check forbids network model downloads")

    DownloadFile._make_http_request = disallow_download
    image = Image.new("RGB", (1000, 160), "white")
    font = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 42)
    ImageDraw.Draw(image).text((30, 40), "环境检查 Knowledge 123", font=font, fill="black")
    cache = Path("D:/dev-tools/cache/knowledgebase-check")
    cache.mkdir(parents=True, exist_ok=True)
    started = time.perf_counter()
    with tempfile.TemporaryDirectory(dir=cache) as directory:
        image_path = Path(directory) / "ocr.png"
        image.save(image_path)
        engine = RapidOCR(config_path=str(model_root / "config.yaml"))
        result = engine(np.asarray(image))
        recognized = " ".join(result.txts or [])
        normalized = recognized.replace(" ", "").lower()
        assert "环境检查" in normalized and "knowledge123" in normalized, recognized
    return recognized, started


# 记录实际环境证据与局限，不将合成检查写成产品验收。
def write_report(checks: list[str], recognized: str, started: float) -> None:
    report = {
        "status": "passed",
        "python": os.sys.version,
        "sqlite": sqlite3.sqlite_version,
        "keyring_backend": type(keyring.get_keyring()).__module__,
        "ocr_text": recognized,
        "ocr_smoke_seconds": round(time.perf_counter() - started, 3),
        "checks": checks,
        "versions": {
            name: importlib.metadata.version(name)
            for name in (
                "rapidocr",
                "onnxruntime",
                "pypdfium2",
                "pdfplumber",
                "pydantic",
                "pyinstaller",
                "pytest",
                "ruff",
            )
        },
        "limitations": "Synthetic environment checks only; no product, real-document quality, Chinese search, provider or installer acceptance.",
    }
    serialized = json.dumps(report, ensure_ascii=False, indent=2) + "\n"
    evidence = (
        Path(__file__).resolve().parent.parent
        / "task-specs/20261020-01-personal-knowledge-agent/environment-check-python.json"
    )
    evidence.write_text(serialized, encoding="utf-8")
    print(serialized)


# 依次运行独立环境检查，报告不代表真实文档或正式接口已验收。
def main() -> None:
    checks = check_data_libraries()
    checks.extend([check_pdf(), check_sqlite(), check_credentials()])
    recognized, started = check_ocr()
    checks.append("RapidOCR CPU offline Chinese/English synthetic image")
    write_report(checks, recognized, started)


if __name__ == "__main__":
    main()
