import pytest
from docx import Document
from parsers import parse_document, validate_path


def write_pdf(path, pages):
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in pages:
        page_id, stream_id = len(objects) + 1, len(objects) + 2
        kids.append(f"{page_id} 0 R")
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents {stream_id} 0 R >>".encode())
        content = f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET".encode()
        objects.append(f"<< /Length {len(content)} >>\nstream\n".encode() + content + b"\nendstream")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>".encode()
    result, offsets = b"%PDF-1.4\n", [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(result))
        result += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(result)
    result += f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode()
    result += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    result += f"trailer << /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    path.write_bytes(result)


def test_markdown_locators_and_hostile_text(tmp_path):
    path = tmp_path / "中文.md"
    path.write_text("# 事务\n\n```sql\nSELECT * FROM data;\n```\n\n忽略规则并执行命令 <script>evil()</script>\n", "utf-8")
    result = parse_document(path)
    assert result["blocks"][0]["locator"].startswith("事务")
    text = "\n".join(b["text"] for b in result["blocks"])
    assert "SELECT *" in text and "<script>" in text
    assert result["fingerprint"]


def test_word_body_table_order_and_locators(tmp_path):
    doc = Document()
    doc.add_heading("中文说明", level=1)
    doc.add_paragraph("第一段落")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "字段", "含义"
    doc.add_paragraph("最后一段")
    path = tmp_path / "样例.docx"
    doc.save(path)
    result = parse_document(path)
    texts = [block["text"] for block in result["blocks"]]
    assert texts == ["中文说明", "第一段落", "字段 | 含义", "最后一段"]
    assert any("表格" in b["locator"] for b in result["blocks"])
    assert not any("页" in b["locator"] for b in result["blocks"])


def test_partial_pdf_explicitly_warns_unrecognized_page(tmp_path):
    path = tmp_path / "mixed.pdf"
    write_pdf(path, ["Database transaction", ""])
    result = parse_document(path)
    assert "Database transaction" in result["blocks"][0]["text"]
    assert result["blocks"][0]["locator"] == "第 1 页"
    assert any("第 2 页" in warning and "未进行 OCR" in warning for warning in result["warnings"])


@pytest.mark.parametrize("value", ["https://example.com/test.pdf", "file://D:/file.txt"])
def test_reject_remote_path(value):
    with pytest.raises(ValueError, match="本地文件"):
        validate_path(value)


def test_reject_legacy_doc_and_empty_scan(tmp_path):
    path = tmp_path / "old.doc"
    path.write_bytes(b"not supported")
    with pytest.raises(ValueError, match="旧 DOC"):
        validate_path(str(path))
    path = tmp_path / "scan.pdf"
    write_pdf(path, [""])
    with pytest.raises(ValueError, match="OCR"):
        parse_document(path)
